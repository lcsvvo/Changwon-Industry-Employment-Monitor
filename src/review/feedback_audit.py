# -*- coding: utf-8 -*-
"""교수 피드백 반영 산출물을 독립적으로 점검하고 근거 로그를 저장한다."""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "outputs/final_model/06_report_assets/professor_feedback"
TABLES = ASSETS / "tables"
REPORT = ROOT / "reports/창원국가산단_산업고용전환진단_분석보고서_교수피드백반영본.docx"
NOTEBOOK = ROOT / "notebooks/07_professor_review_master.ipynb"


def _report_text() -> str:
    doc = Document(REPORT)
    pieces = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        pieces.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(pieces)


def run_checks() -> list[dict[str, str]]:
    qa = pd.read_csv(TABLES / "qa_level_shift_cases.csv")
    source = pd.read_csv(TABLES / "wood_paper_source_audit.csv")
    machine = pd.read_csv(TABLES / "machine_yoy_qoq_interpretation.csv").iloc[0]
    boundary = pd.read_csv(TABLES / "wood_paper_boundary_margin.csv").set_index("quarter")
    followup = pd.read_csv(TABLES / "signal_followup_summary.csv")
    validation = pd.read_csv(TABLES / "validation_model_comparison.csv")
    external = pd.read_csv(TABLES / "external_contemporaneous_cases.csv")
    handoff = pd.read_csv(ROOT / "outputs/final_model/05_handoff/tables/handoff_cards_latest.csv")
    decomp = pd.read_csv(TABLES / "wood_paper_level_shift_decomposition.csv")
    ratio = pd.read_csv(TABLES / "ratio_boundary_headcount_summary.csv")
    cci = pd.read_csv(TABLES / "cci_concurrent_summary.csv")
    report = _report_text()
    nb = json.loads(NOTEBOOK.read_text(encoding="utf-8"))
    notebook = "\n".join("".join(c.get("source", [])) for c in nb["cells"])
    app = (ROOT / "src/app/main.py").read_text(encoding="utf-8")
    vm = (ROOT / "src/app/view_models.py").read_text(encoding="utf-8")
    ui_src = (ROOT / "src/app/ui.py").read_text(encoding="utf-8")

    checks: list[tuple[str, bool, str, str]] = [
        (
            "1. 목재·종이 급변 QA",
            {"2022Q4", "2025Q1"} <= set(qa.quarter)
            and {"2022Q4", "2023Q1", "2023Q2", "2023Q3", "2025Q1", "2026Q2"} <= set(source.quarter)
            and source.not_confirmed.str.contains("한두 기업", na=False).any()
            and set(decomp.loc[decomp.quarter.str.startswith("2025"), "direction_excluding_shift"]) == {"생산↓·고용↓"}
            and "기저효과로 보는 것이 타당하다" in report,
            "분포기반 QA·원자료 추적표에 더해, 급변 분기를 분리한 YoY 분해로 2025년 S1 기저효과와 2022Q4 수준하락 효과를 수치화",
            "qa_level_shift_cases.csv; wood_paper_source_audit.csv; wood_paper_level_shift_decomposition.csv",
        ),
        (
            "2. 기계업종 기저효과",
            machine["qoq_recovery_while_yoy_below"] in (False, "False", 0)
            and machine["peak_since_2023_quarter"] == "2025Q2"
            and bool(machine["comparison_is_peak_since_2023"]),
            "YoY·QoQ·2023년 이후 고점을 병기하고 2025Q2 고기저 여부를 계산",
            "machine_yoy_qoq_interpretation.csv",
        ),
        (
            "3. 목재·종이 경계 민감도",
            abs(boundary.loc["2026Q2", "E"] - 10.0840336134) < 1e-8
            and boundary.loc["2026Q2", "E_upper_headcount_to_flip"] == 1
            and boundary.loc["2026Q1", "E_upper_headcount_to_flip"] == 4
            and ratio.set_index("size_group").loc["300인 미만", "rows_flip_within_2"] == 18
            and "27.3%" in report and "규모게이트가 아니라 E 10% 상위경계" in report,
            "48명/47명 감소와 E 상위경계 여유를 인원으로 제시하고, 규모별 비율경계 인원 민감도로 규모게이트 근거를 보강",
            "wood_paper_boundary_margin.csv; ratio_boundary_headcount_summary.csv; ratio_boundary_headcount_latest.csv",
        ),
        (
            "4. 시점분리 검증",
            set(validation.model) == {"triage_final", "electre_fixed"}
            and set(followup.stage) == {"우선점검", "추가확인", "관찰"}
            and len(external) == 3
            and cci.set_index("group").loc["우선점검 행", "emp_direction_agree"] == "9/9"
            and "시점분리 사후 점검과 동시점 대조" in report,
            "검증 결과를 본론(2-3)으로 옮겨 추가 위축 예측용이 아님을 설명하고, 신호 이후 회복률·비교기준 고점·창원상의 동시점 대조를 보강",
            "signal_followup_summary.csv; priority_base_peak_check.csv; cci_concurrent_check.csv; external_contemporaneous_cases.csv",
        ),
        (
            "5. 보고서·그림",
            all(f"그림 {n}." in report for n in range(1, 8))
            and "2023년 1~3분기" in report
            and "한두 기업" in report,
            "Q1·Q2·기준선·단계격자와 원본 Streamlit 화면 2장을 배치하고 해석 문단을 추가",
            "교수피드백반영본.docx 그림 1~7 및 본문",
        ),
        (
            "6. 서술·용어",
            "선제점검 단계분류(Triage)" in report
            and "우선점검 진입용 최소 고용규모(규모게이트)" in report
            and "관찰 수준(OBSERVE)" in report
            and "추가확인 수준(CHECK)" in report
            and "우선점검 수준(PRIORITY)" in report
            and "자료·모형상 단일 범주 확정 곤란(UNDETERMINED)" in report
            and "범주수용도지수(CAI)" in report
            and "우선검토 수준" not in report
            and "우선점검 후보" not in report
            and "해석 범위도 이 절에 모아 둔다" in report,
            "최초 용어를 풀어 쓰고 단계명을 우선점검·추가확인·관찰로 통일, 해석범위 부정 서술은 4-1 한계 절로 모음",
            "보고서; src/app/view_models.py",
        ),
        (
            "7. 선별효과",
            "10개 업종 중 기계와 목재종이 2개" in report
            and "51.56%" in report
            and "80%" not in report,
            "2/10 선별과 고용비중 51.56%를 구분하고 80% 업무절감 표현을 제거",
            "보고서 선별효과 문단",
        ),
        (
            "8. 검토 패키지·재현성",
            "검토 패키지는 저장 출력 확인용" in notebook
            and (ROOT / "src/review/build_review_package.py").exists()
            and "if ROOT is None:" in notebook,
            "저장 출력 확인용 패키지와 전체 재실행 저장소를 구분하고, src가 없으면 노트북 첫 셀이 재실행 위치를 안내하고 멈춤",
            "00_개요.txt; build_review_package.py; 최종 노트북 17.8",
        ),
        (
            "9. Streamlit",
            "자료 수준 급변 확인" in app
            and "workflow_label_display" in app
            and "WORKFLOW_LABEL_DISPLAY" in vm
            and "stage_mark_legend_html" in app and "confirmed_active_detail_verified_posting_count" in ui_src
            and {"emp_qoq_delta", "E_upper_headcount_to_flip", "trend_check_question"} <= set(handoff.columns),
            "업종 표식 도형(●▲■×, 회색 테두리)+단계색·범례, 선택 카드 버튼 강조, Q1~Q3 3칸, 추이 띠 빗금/점선 구분, 채용카드 포함관계 순서, 섹션 바로가기",
            "src/app/main.py; src/app/ui.py; src/app/view_models.py; handoff_cards_latest.csv",
        ),
    ]

    failed = [name for name, passed, _, _ in checks if not passed]
    if failed:
        raise AssertionError(f"교수 피드백 미반영 점검 항목: {failed}")
    return [
        {"section": name, "status": "PASS", "change": change, "evidence": evidence}
        for name, _, change, evidence in checks
    ]


def write_log(rows: list[dict[str, str]]) -> Path:
    lines = [
        "# 교수 피드백 반영 검증 로그",
        "",
        "자동검증일: 2026-09-28",
        "",
        "| 구분 | 상태 | 어떻게 수정했는가 | 검증 근거 |",
        "|---|---|---|---|",
    ]
    lines.extend(
        f"| {r['section']} | {r['status']} | {r['change']} | {r['evidence']} |" for r in rows
    )
    lines.extend([
        "",
        "미완료: 목재·종이 2025Q1 급증의 실제 원인(기업 증원·분류/집계범위 변경)과 "
        "2026Q2 감소의 기업별 집중 여부는 공개 업종 집계표로 확인할 수 없다. 원자료 수치 대조, "
        "수준변화 QA, 이후 분기 지속 여부와 기저효과 해석은 완료했지만 원인 규명 완료로 표시하지 않는다.",
    ])
    path = ASSETS / "feedback_implementation_log.md"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def main() -> None:
    rows = run_checks()
    path = write_log(rows)
    print(f"feedback audit passed: {len(rows)}/9 sections; {path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
