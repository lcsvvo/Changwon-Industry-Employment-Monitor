# -*- coding: utf-8 -*-
"""관세청 HS부호 원본 → 유효 HS6 모집단 + KICOX 업종 대응.

입력 (둘 다 관세청 공표 원본, data/raw/customs/reference/)
    customs_hs_code_20260101.xlsx                 HSK 10단위 + 적용기간
    customs_hsk_nature_classification_20260101.xlsx  세번 2/4/6단위 품명

출력
    data/processed/final_model/reference/industry_crosswalk/hs6_universe_customs.csv

매핑 근거 — 무엇을 쓰고 무엇을 쓰지 않았나
    **쓴 것**: HS 류(2단위)와 관세청이 공표한 `세번2단위품명`.
      류는 국제 표준이고 품명은 관세청 원본에 그대로 실려 있어 재현 가능하다.

    **쓰지 않은 것**: 관세청 '신성질별 분류'(대분류 1.소비재 / 2.원자재 / 3.자본재).
      이것은 **용도별** 분류이지 생산 업종 분류가 아니다. KICOX 10업종과 축이 다르므로
      업종 대응 근거로 삼지 않았다. 다만 참고용으로 열에 보존한다.

    그래서 대응은 류 단위에서만, 아래 세 등급으로만 만든다.
      A  류 전체가 한 KICOX 업종에 대응 (72·73 철강, 87 운송장비, 89 운송장비)
      B  류의 대부분이 한 업종이나 다른 업종이 섞임 (84 기계, 85 전기전자)
      제외  한 업종으로 귀속할 근거가 없는 류 (74~81 비철, 90 광학·의료 등)

    B 등급은 '확인신호' 가 아니라 맥락으로만 쓴다. 제외한 류는 아예 수집하지 않는다.

실행:  python src/evidence/build_hs6_universe.py
"""
from __future__ import annotations

import csv
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
REF = ROOT / "data/raw/customs/reference"
HS_FILE = REF / "customs_hs_code_20260101.xlsx"
NATURE_FILE = REF / "customs_hsk_nature_classification_20260101.xlsx"
OUT = ROOT / "data/processed/final_model/reference/industry_crosswalk/hs6_universe_customs.csv"

VALID_FROM_YEAR = 2026          # 이 연도에도 적용되는 코드만 유효로 본다

# (류, KICOX 업종, 등급, 근거). 여기 없는 류는 수집 대상에서 제외한다.
CHAPTER_MAP = {
    "72": ("철강", "A", "제72류 철강. KICOX '철강'에 직접 대응"),
    "73": ("철강", "A", "제73류 철강의 제품. KICOX '철강'에 직접 대응"),
    "84": ("기계", "B", "제84류 원자로·보일러·기계류. 전산기기(8471 등)가 섞여 있어 "
                        "류 전체가 '기계'는 아니다 — 맥락 용도로만"),
    "85": ("전기전자", "B", "제85류 전기기기·음향영상기기. KICOX '전기전자'가 중심이나 "
                            "일부 부품이 다른 업종과 겹친다"),
    "87": ("운송장비", "A", "제87류 철도 외 차량과 부품. KICOX '운송장비'에 직접 대응"),
    "89": ("운송장비", "A", "제89류 선박. KICOX '운송장비'(조선)에 직접 대응"),
}
EXCLUDED_NOTE = {
    "74-81": "비철금속. KICOX '철강'이 비철을 포함하는지 확정 불가",
    "90": "광학·의료·정밀기기. KSIC 27 대응 여부가 불확실한 것과 같은 이유",
    "기타": "한 KICOX 업종으로 귀속할 공식 근거가 없음",
}
COLS = ["hs6", "hs_chapter", "chapter_name_ko", "hs6_name_ko",
        "target_kicox_industry", "mapping_type", "mapping_confidence",
        "rationale", "nature_major_code", "nature_major_name",
        "valid_through", "source"]
SRC = ("관세청_HS부호(공공데이터포털 15049722, 2026-01-01 기준) + "
       "관세청_HSK별 신성질별_성질별 분류(15049720)")


def load_hs_codes() -> dict[str, dict]:
    """HS6 → {유효성, 성질코드}. HSK 10단위에서 앞 6자리를 취한다."""
    wb = openpyxl.load_workbook(HS_FILE, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    iH, iEnd = hdr.index("HS부호"), hdr.index("적용종료일자")
    iNc, iNn = hdr.index("성질통합분류코드"), hdr.index("성질통합분류코드명")
    out: dict[str, dict] = {}
    for r in it:
        code = str(r[iH] or "").strip()
        if len(code) < 6 or not code.isdigit():
            continue
        end = r[iEnd]
        year = getattr(end, "year", None)
        if year is not None and year < VALID_FROM_YEAR:
            continue                                   # 적용 종료된 코드는 제외
        h6 = code[:6]
        prev = out.get(h6)
        if prev is None:
            out[h6] = dict(nature_code=r[iNc], nature_name=r[iNn],
                           valid_through=year)
        elif year and (prev["valid_through"] or 0) < year:
            prev["valid_through"] = year
    wb.close()
    return out


def load_names() -> tuple[dict[str, str], dict[str, str]]:
    """(류 품명, HS6 품명). 관세청 공표 품명을 그대로 쓴다."""
    wb = openpyxl.load_workbook(NATURE_FILE, read_only=True, data_only=True)
    sheet = wb.sheetnames[-1]                          # 가장 최근 연도 시트
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    # 시트마다 첫 열 이름이 다르다(2025년='HS10단위부호', 2026년='국제적 상품분류체계(HS)10단위부호').
    iH = next(i for i, h in enumerate(hdr) if h and "10단위부호" in str(h))
    i2, i6 = hdr.index("세번2단위품명"), hdr.index("세번6단위품명")
    chapters: dict[str, str] = {}
    names: dict[str, str] = {}
    for r in it:
        code = str(r[iH] or "").strip()
        if len(code) < 6 or not code.isdigit():
            continue
        ch, h6 = code[:2], code[:6]
        if ch not in chapters and r[i2]:
            chapters[ch] = str(r[i2]).strip()
        if h6 not in names and r[i6]:
            names[h6] = str(r[i6]).strip()
    wb.close()
    return chapters, names


def main() -> int:
    if not HS_FILE.exists() or not NATURE_FILE.exists():
        print("[skip] 관세청 기준자료 없음 — collect_customs_hs_reference.py 먼저 실행")
        return 0
    codes = load_hs_codes()
    chapters, names = load_names()
    rows = []
    for h6, info in sorted(codes.items()):
        ch = h6[:2]
        if ch not in CHAPTER_MAP:
            continue                                   # 근거 없는 류는 아예 담지 않는다
        industry, grade, why = CHAPTER_MAP[ch]
        rows.append({
            "hs6": h6, "hs_chapter": ch,
            "chapter_name_ko": chapters.get(ch, ""),
            "hs6_name_ko": names.get(h6, ""),
            "target_kicox_industry": industry,
            "mapping_type": "direct" if grade == "A" else "approximate",
            "mapping_confidence": grade, "rationale": why,
            "nature_major_code": info["nature_code"] or "",
            "nature_major_name": info["nature_name"] or "",
            "valid_through": info["valid_through"] or "",
            "source": SRC,
        })
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLS)
        w.writeheader()
        w.writerows(rows)

    total_valid = len(codes)
    print(f"[ok] {OUT.relative_to(ROOT)}")
    print(f"     전체 유효 HS6 {total_valid:,}개 중 근거 있는 류만 {len(rows):,}개 수록")
    by_ch: dict[str, int] = {}
    for r in rows:
        by_ch[r["hs_chapter"]] = by_ch.get(r["hs_chapter"], 0) + 1
    for ch in sorted(by_ch):
        ind, grade, _ = CHAPTER_MAP[ch]
        print(f"     {ch}류 {chapters.get(ch, '')[:22]:<24} {by_ch[ch]:4d}개 "
              f"→ {ind} (등급 {grade})")
    print(f"     제외한 류: {', '.join(f'{k}({v[:24]})' for k, v in EXCLUDED_NOTE.items())}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
