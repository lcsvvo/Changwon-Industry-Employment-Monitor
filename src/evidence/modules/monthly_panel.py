# -*- coding: utf-8 -*-
"""월별 KICOX 업종 원자료 → 분기 신호 지속성 보조정보.

무엇을 하는가
    분기 판정과 같은 방향의 신호가 그 분기의 3개월 중 몇 개월에서 관측되는지 센다.

무엇을 하지 않는가
    - Triage stage 를 바꾸는 새 조건을 만들지 않는다.
    - 월별 자료를 쓴다고 해서 독립관측이 3배가 되는 것이 아니다. 분기말 단일시점
      의존을 줄이는 보조정보일 뿐이다.
    - 미래 월을 참조하지 않는다. 분기 t 의 지속성은 분기 t 에 속한 달과
      그 12개월 전 달만 사용한다.

자료 출처와 한계
    생산  : 공공데이터포털 월별 원자료(2021M01~2024M03) + 연간보정본 월별(2022M10~2024M03)
    고용/가동업체/가동률 : 연간보정본 월별(2022M10~2024M03) 만 존재
    2024Q2 부터 KICOX 공표주기가 분기로 바뀌어 월별 자료 자체가 없다.
    따라서 지속성은 일부 분기에서만 계산 가능하며, 나머지는 결측으로 남긴다.
"""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path

import numpy as np
import pandas as pd

try:
    import openpyxl
except ImportError:                                   # pragma: no cover
    openpyxl = None

from .config import INDUSTRIES

COMPLEX = "창원"
PORTAL_DATASETS = {"production": "15085898"}          # 업종별 월별이 존재하는 포털 데이터셋
PORTAL_QUARTERLY_FROM = "20240401"                    # 이 날짜 이후 포털 공표본은 분기값
REV_SHEETS = {"production": "표5", "employment": "표9",
              "firms_op": "표3", "op_rate": "표12"}
# stock/ratio 는 월말 시점값, flow 는 그 달의 합계
KIND = {"production": "flow", "employment": "stock",
        "firms_op": "stock", "op_rate": "ratio"}


def _to_num(v):
    s = str(v).strip().replace(",", "") if v is not None else ""
    if s.upper() in ("X", "×") or s in ("", "-", "N/A", "None", "nan"):
        return np.nan
    try:
        return float(s)
    except ValueError:
        return np.nan


def _read_rows(path: Path):
    raw = path.read_bytes()
    for enc in ("cp949", "utf-8-sig", "utf-8"):
        try:
            return list(csv.reader(io.StringIO(raw.decode(enc))))
        except UnicodeDecodeError:
            continue
    return list(csv.reader(io.StringIO(raw.decode("cp949", "replace"))))


def _clean(h) -> str:
    return str(h).split("(")[0].strip() if h is not None else ""


def load_portal_monthly(root: Path) -> pd.DataFrame:
    """포털 월별 CSV 에서 창원 행의 업종별 값을 읽는다 (원자료 vintage)."""
    d = root / "data" / "raw" / "kicox" / "core"
    recs = []
    for indicator, pk in PORTAL_DATASETS.items():
        for fp in sorted(d.glob(f"{pk}_*.csv")):
            date = fp.name.split("_")[1][:8]
            if date >= PORTAL_QUARTERLY_FROM:
                # 2024Q2 부터 KICOX 공표주기가 분기다. 이 시점 이후의 포털 파일은
                # 분기값이므로 월값으로 취급하면 안 된다(월 YoY 정의가 깨진다).
                continue
            rows = _read_rows(fp)
            if not rows:
                continue
            hdr = [_clean(h) for h in rows[0]]
            cw = next((r for r in rows[1:] if r and r[0].strip() == COMPLEX), None)
            if cw is None:
                continue
            for i, h in enumerate(hdr):
                if h in INDUSTRIES and i < len(cw):
                    recs.append(dict(indicator=indicator, industry=h,
                                     month=f"{date[:4]}-{date[4:6]}",
                                     value=_to_num(cw[i]), vintage="portal_raw"))
    return pd.DataFrame(recs)


def load_revision_monthly(root: Path) -> pd.DataFrame:
    """연간보정본 월별 xlsx 에서 창원 행의 업종별 값을 읽는다 (revised vintage)."""
    if openpyxl is None:
        return pd.DataFrame(columns=["indicator", "industry", "month", "value", "vintage"])
    recs = []
    for fp in sorted((root / "data" / "raw" / "kicox" / "revision").rglob("*.xlsx")):
        m = re.match(r"^(\d{4})M(\d{1,2})$", fp.stem)
        if not m:
            continue                                   # 분기 보정본은 월별 지속성에 쓰지 않는다
        month = f"{m.group(1)}-{int(m.group(2)):02d}"
        wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
        for indicator, prefix in REV_SHEETS.items():
            sheet = next((s for s in wb.sheetnames if s.startswith(prefix + " ")), None)
            if sheet is None:
                continue
            ws = wb[sheet]
            hdr = cw = None
            for row in ws.iter_rows(values_only=True):
                first = str(row[0]).strip() if row and row[0] is not None else ""
                if first == "산업단지":
                    hdr = [_clean(c) for c in row]
                elif first == COMPLEX and hdr is not None:
                    cw = row
                    break
            if hdr is None or cw is None:
                continue
            for i, h in enumerate(hdr):
                if h in INDUSTRIES and i < len(cw):
                    recs.append(dict(indicator=indicator, industry=h, month=month,
                                     value=_to_num(cw[i]), vintage="annual_revision"))
        wb.close()
    return pd.DataFrame(recs)


def build_monthly_panel(root: Path) -> pd.DataFrame:
    """월 × 업종 × 지표 패널. 같은 월에 둘 다 있으면 보정본을 채택한다.

    보정본 값이 비공개(X, 조사대상 5개사 이하)여서 결측인 달도 보정본을 그대로
    채택한다. 포털 원자료 값으로 되메우면 마스터(`build_changwon_master.py`)의
    결측정책과 어긋나 같은 분기가 두 값을 갖게 되기 때문이다.
    """
    df = pd.concat([load_portal_monthly(root), load_revision_monthly(root)],
                   ignore_index=True)
    if df.empty:
        return df
    order = {"annual_revision": 0, "portal_raw": 1}
    df["_p"] = df["vintage"].map(order)
    df = (df.sort_values(["indicator", "industry", "month", "_p"])
            .drop_duplicates(["indicator", "industry", "month"], keep="first")
            .drop(columns="_p").reset_index(drop=True))
    df["quarter"] = df["month"].str[:4] + "Q" + (
        ((df["month"].str[5:7].astype(int) - 1) // 3 + 1).astype(str))
    return df


def add_monthly_yoy(panel: pd.DataFrame) -> pd.DataFrame:
    """월별 전년동월 대비. 분기 YoY 정의(전년 동일 달력시점 대비)와 같은 방식이며
    결측월을 건너뛰어 붙이지 않는다."""
    if panel.empty:
        return panel
    d = panel.copy()
    lag = d[["indicator", "industry", "month", "value", "vintage"]].copy()
    lag["month"] = (pd.PeriodIndex(lag["month"], freq="M") + 12).astype(str)
    lag = lag.rename(columns={"value": "value_lag12", "vintage": "vintage_lag12"})
    d = d.merge(lag, on=["indicator", "industry", "month"], how="left", validate="one_to_one")
    d["yoy_pct"] = (d["value"] / d["value_lag12"] - 1) * 100
    d.loc[~np.isfinite(d["value_lag12"]) | (d["value_lag12"] == 0), "yoy_pct"] = np.nan
    d["yoy_pct"] = d["yoy_pct"].replace([np.inf, -np.inf], np.nan)
    d["vintage_consistent"] = (d["vintage"] == d["vintage_lag12"]).where(
        d["vintage_lag12"].notna())
    return d


def quarterly_persistence(monthly_yoy: pd.DataFrame) -> pd.DataFrame:
    """분기별 '음(-)의 월 수 / 관측 가능한 월 수'.

    months_observed < 3 이면 3개월 중 몇 개월인지 단정하지 않는다.
    """
    if monthly_yoy.empty:
        return pd.DataFrame(columns=["industry", "quarter"])
    d = monthly_yoy.dropna(subset=["yoy_pct"]).copy()
    d["neg"] = d["yoy_pct"] < 0
    g = (d.groupby(["indicator", "industry", "quarter"])
           .agg(months_observed=("yoy_pct", "size"),
                negative_months=("neg", "sum"),
                min_yoy=("yoy_pct", "min"),
                max_yoy=("yoy_pct", "max"),
                vintage_consistent_months=("vintage_consistent", "sum"))
           .reset_index())
    out = None
    for ind in g["indicator"].unique():
        sub = g[g["indicator"] == ind].drop(columns="indicator")
        sub = sub.rename(columns={c: f"{ind}_{c}" for c in sub.columns
                                  if c not in ("industry", "quarter")})
        out = sub if out is None else out.merge(sub, on=["industry", "quarter"], how="outer")
    return out.sort_values(["industry", "quarter"]).reset_index(drop=True)
