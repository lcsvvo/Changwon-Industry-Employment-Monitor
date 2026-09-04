# -*- coding: utf-8 -*-
"""
qa_master.py
창원국가산업단지 KICOX 산업동향 마스터 — 전처리 품질 QA (읽기 전용)

이 스크립트는 data/processed/*.csv 를 절대 수정하지 않는다.
검사 항목 A~H 를 실행하고 PASS/FAIL 과 실제 수치를 stdout 및
logs/qa_report.txt 에 기록한다. 실패 항목이 하나라도 있으면 exit code 1.

실행
    python src/qa_master.py
"""

from __future__ import annotations

import os
import re
import sys
from datetime import datetime

import numpy as np
import pandas as pd

try:
    import openpyxl
except ImportError:
    openpyxl = None

# Windows 콘솔의 기본 cp949에서도 한글과 특수문자 로그가 깨지거나
# UnicodeEncodeError로 중단되지 않도록 출력 인코딩을 고정한다.
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIR_PROC = os.path.join(BASE, "data", "processed")
DIR_REV = os.path.join(BASE, "data", "annual_revision")
DIR_LOG = os.path.join(BASE, "logs")

P_IND = os.path.join(DIR_PROC, "changwon_industry_master.csv")
P_TOT = os.path.join(DIR_PROC, "changwon_total_master.csv")

INDUSTRIES = ["음식료", "섬유의복", "목재종이", "석유화학",
              "비금속", "철강", "기계", "전기전자", "운송장비", "기타"]

NUMERIC_IND_COLS = ["production", "employment", "op_rate_official", "op_rate_approx",
                     "firms_in", "firms_op"]
NUMERIC_TOT_COLS = ["production_total", "employment_total",
                     "firms_in_total", "op_rate_total"]

SOURCE_COLS = ["production_source", "employment_source", "op_rate_source",
               "firms_in_source", "firms_op_source"]
ALLOWED_SOURCES = {"data_portal_monthly", "data_portal_quarterly",
                    "kicox_annual_revision", "kicox_republication"}

STR_JUNK = {"X", "-", "", "null", "nan", "N/A", "None"}

LOG_LINES: list[str] = []
RESULTS: list[tuple[str, str, bool]] = []   # (section, label, passed)


def log(msg: str = "") -> None:
    print(msg)
    LOG_LINES.append(str(msg))


def record(section: str, label: str, passed: bool) -> bool:
    RESULTS.append((section, label, passed))
    tag = "PASS" if passed else "FAIL"
    log(f"  [{tag}] {label}")
    return passed


def hr(title: str) -> None:
    log("\n" + "=" * 78)
    log(title)
    log("=" * 78)


# ----------------------------------------------------------------------------
# 유틸: 원본 CSV를 문자열로 다시 읽기 (숫자 컬럼 잔재 검사용)
# ----------------------------------------------------------------------------

def _read_str(path: str) -> pd.DataFrame:
    return pd.read_csv(path, dtype=str, keep_default_na=False, na_values=[])


# ----------------------------------------------------------------------------
# A. 구조
# ----------------------------------------------------------------------------

def check_a(ind: pd.DataFrame, tot: pd.DataFrame) -> None:
    hr("A. 구조")

    record("A", f"industry master 행수 = 340 (실제 {len(ind)})", len(ind) == 340)
    n_q = ind.quarter.nunique()
    record("A", f"quarter 34개 (실제 {n_q})", n_q == 34)
    n_i = ind.industry.nunique()
    record("A", f"industry 10개 (실제 {n_i})", n_i == 10)

    record("A", f"total master 행수 = 34 (실제 {len(tot)})", len(tot) == 34)

    dup_ind = ind.duplicated(subset=["quarter", "industry"]).sum()
    record("A", f"(quarter, industry) 중복 0건 (실제 {dup_ind})", dup_ind == 0)

    dup_tot = tot.duplicated(subset=["quarter"]).sum()
    record("A", f"total quarter 중복 0건 (실제 {dup_tot})", dup_tot == 0)

    expected_set = set(INDUSTRIES)
    actual_set = set(ind.industry.unique())
    diff = expected_set.symmetric_difference(actual_set)
    record("A", f"업종명 집합 정확히 일치 (차이: {diff or '없음'})", len(diff) == 0)

    per_q = ind.groupby("quarter").industry.nunique()
    bad_q = per_q[per_q != 10]
    record("A", f"34개 분기 전체 업종 수 = 10 (위반 분기: {bad_q.to_dict() or '없음'})",
           len(bad_q) == 0)

    qfmt_ok = ind.quarter.astype(str).str.match(r"^\d{4}Q[1-4]$")
    bad_fmt = sorted(ind.loc[~qfmt_ok, "quarter"].unique())
    record("A", f"quarter 형식 YYYYQ[1-4] 전부 일치 (위반: {bad_fmt or '없음'})",
           qfmt_ok.all())

    qs = sorted(ind.quarter.unique())
    if qs:
        y0, k0 = int(qs[0][:4]), int(qs[0][-1])
        y1, k1 = int(qs[-1][:4]), int(qs[-1][-1])
        expected_seq = []
        y, k = y0, k0
        while (y, k) <= (y1, k1):
            expected_seq.append(f"{y}Q{k}")
            k += 1
            if k == 5:
                k = 1
                y += 1
        missing = [q for q in expected_seq if q not in qs]
        log(f"  범위 : {qs[0]} ~ {qs[-1]} (연속 기대 {len(expected_seq)}개, 실제 {len(qs)}개)")
        record("A", f"2018Q1~2026Q2 연속, 빠진 분기 없음 (누락: {missing or '없음'})",
               len(missing) == 0 and qs[0] == "2018Q1" and qs[-1] == "2026Q2")
    else:
        record("A", "quarter 데이터 없음", False)


# ----------------------------------------------------------------------------
# B. 데이터 타입
# ----------------------------------------------------------------------------

def check_b(ind: pd.DataFrame, tot: pd.DataFrame) -> None:
    hr("B. 데이터 타입")

    for c in NUMERIC_IND_COLS:
        is_num = pd.api.types.is_numeric_dtype(ind[c])
        record("B", f"industry.{c} numeric dtype (실제 {ind[c].dtype})", is_num)

    for c in NUMERIC_TOT_COLS:
        is_num = pd.api.types.is_numeric_dtype(tot[c])
        record("B", f"total.{c} numeric dtype (실제 {tot[c].dtype})", is_num)

    ind_str = _read_str(P_IND)
    tot_str = _read_str(P_TOT)
    junk_found = {}
    for c in NUMERIC_IND_COLS:
        bad = ind_str[c].isin(STR_JUNK - {""}) | (ind_str[c] == "")
        # 원본에 빈칸('')은 정상 결측(NaN)으로 pandas가 읽으므로 문자열 잔재 검사는
        # 'X','-','null','nan' 등 실제 텍스트 잔재만 본다. 빈칸은 별도 카운트.
        literal_bad = ind_str[c].isin({"X", "-", "null", "nan"})
        if literal_bad.any():
            junk_found[c] = int(literal_bad.sum())
    for c in NUMERIC_TOT_COLS:
        literal_bad = tot_str[c].isin({"X", "-", "null", "nan"})
        if literal_bad.any():
            junk_found[c] = int(literal_bad.sum())
    record("B", f"숫자 컬럼에 문자열 잔재('X','-','null','nan') 없음 (발견: {junk_found or '없음'})",
           len(junk_found) == 0)


# ----------------------------------------------------------------------------
# C. 결측
# ----------------------------------------------------------------------------

def check_c(ind: pd.DataFrame, tot: pd.DataFrame) -> None:
    hr("C. 결측")

    prod_na_q = sorted(ind.loc[ind.production.isna(), "quarter"].unique())
    log(f"  production 결측 분기: {prod_na_q}")
    record("C", "의도된 결측: industry production = 2023Q4 10행만",
           prod_na_q == ["2023Q4"] and int(ind.production.isna().sum()) == 10)

    emp_na = int(ind.employment.isna().sum())
    record("C", f"employment 결측 0 (실제 {emp_na})", emp_na == 0)

    firms_in_na = int(ind.firms_in.isna().sum())
    firms_op_na = int(ind.firms_op.isna().sum())
    record("C", f"firms_in 결측 0 (실제 {firms_in_na})", firms_in_na == 0)
    record("C", f"firms_op 결측 0 (실제 {firms_op_na})", firms_op_na == 0)

    off_na_q = sorted(ind.loc[ind.op_rate_official.isna(), "quarter"].unique())
    app_na_q = sorted(ind.loc[ind.op_rate_approx.isna(), "quarter"].unique())
    log(f"  op_rate_official 결측 분기 ({len(off_na_q)}개): "
        f"{off_na_q if len(off_na_q) <= 12 else str(off_na_q[:12]) + '...'}")
    log(f"  op_rate_approx  결측 분기 ({len(app_na_q)}개): "
        f"{app_na_q if len(app_na_q) <= 12 else str(app_na_q[:12]) + '...'}")
    both_filled = ind.op_rate_official.notna() & ind.op_rate_approx.notna()
    n_both = int(both_filled.sum())
    record("C", f"op_rate_official/approx 동시 채워진 행 0건 (실제 {n_both})", n_both == 0)

    tot_na = {c: int(tot[c].isna().sum()) for c in NUMERIC_TOT_COLS}
    record("C", f"total master 4개 지표 결측 0 (실제 {tot_na})",
           all(v == 0 for v in tot_na.values()))

    # masked 플래그 일관성: masked==1 인 셀은 값이 NaN 이어야 한다
    mask_incons = {}
    for c in ["production", "employment", "firms_in", "firms_op"]:
        mcol = f"{c}_masked"
        bad = ind[(ind[mcol] == 1) & ind[c].notna()]
        if len(bad):
            mask_incons[c] = len(bad)
    both_op_na = ind.op_rate_official.isna() & ind.op_rate_approx.isna()
    op_bad = ind[(ind.op_rate_masked == 1) & ~both_op_na]
    if len(op_bad):
        mask_incons["op_rate"] = len(op_bad)
    record("C", f"_masked==1 셀은 모두 NaN (역방향 불일치: {mask_incons or '없음'})",
           len(mask_incons) == 0)

    prod_2023q4_masked = ind.loc[ind.quarter == "2023Q4", "production_masked"]
    record("C", f"production 2023Q4 10행 모두 production_masked==1 "
                f"(실제 {int((prod_2023q4_masked == 1).sum())}/10)",
           (prod_2023q4_masked == 1).all() and len(prod_2023q4_masked) == 10)

    # 참고용: masked==0 인데 값이 NaN인 셀 (마스킹과 무관한 결측, 실패 아님/정보성)
    info_gaps = {}
    for c in ["production", "employment", "firms_in", "firms_op"]:
        mcol = f"{c}_masked"
        gap = ind[(ind[c].isna()) & (ind[mcol] == 0)]
        if len(gap):
            info_gaps[c] = sorted(gap.quarter.unique())
    op_gap = ind[both_op_na & (ind.op_rate_masked == 0)]
    if len(op_gap):
        info_gaps["op_rate(both)"] = sorted(op_gap.quarter.unique())
    log(f"  [정보] masked==0 인데 값이 NaN인 셀(비-마스킹 결측, 실패 아님): {info_gaps or '없음'}")


# ----------------------------------------------------------------------------
# D. 값 범위 · 단위 일관성
# ----------------------------------------------------------------------------

def check_d(ind: pd.DataFrame, tot: pd.DataFrame) -> None:
    hr("D. 값 범위 · 단위 일관성")

    prod_notna = ind[ind.production.notna()]
    neg_prod = int((prod_notna.production < 0).sum())
    zero_prod = prod_notna[prod_notna.production == 0]
    if len(zero_prod):
        log(f"  [정보] production == 0 인 셀 {len(zero_prod)}건 (음수는 아님, 실제 소규모 업종 값으로 보임): "
            f"{sorted(zero_prod.quarter.unique())} / 업종 {sorted(zero_prod.industry.unique())}")
    record("D", f"production 음수 0건 (2023Q4 NaN 제외, 실제 음수 {neg_prod}건)", neg_prod == 0)

    neg_emp = int((ind.employment < 0).sum())
    record("D", f"employment >= 0 (위반 {neg_emp}건)", neg_emp == 0)

    op_combined = ind.op_rate_official.combine_first(ind.op_rate_approx)
    op_valid = op_combined.dropna()
    out_of_range = int(((op_valid < 0) | (op_valid > 100)).sum())
    tot_out = int(((tot.op_rate_total < 0) | (tot.op_rate_total > 100)).sum())
    record("D", f"op_rate 값이 0~100 (industry 위반 {out_of_range}건, total 위반 {tot_out}건)",
           out_of_range == 0 and tot_out == 0)

    viol = ind[ind.firms_in < ind.firms_op]
    if len(viol):
        log("  firms_in < firms_op 위반 행:")
        log(viol[["quarter", "industry", "firms_in", "firms_op"]].to_string(index=False))
    record("D", f"firms_in >= firms_op (위반 {len(viol)}건)", len(viol) == 0)

    # 단위 급변 탐지 (제거하지 않음, 발견만 보고)
    log("\n  단위 급변 탐지 (분기별 합계, 인접 분기 대비 >=5x 또는 <=1/5)")
    jump_found = []
    for c in ["production", "employment", "firms_in", "firms_op"]:
        s = ind.groupby("quarter")[c].sum(min_count=1).sort_index()
        s_valid = s.dropna()
        qs = list(s_valid.index)
        for i in range(1, len(qs)):
            prev, cur = s_valid.iloc[i - 1], s_valid.iloc[i]
            if prev == 0 or cur == 0:
                continue
            ratio = cur / prev
            if ratio >= 5 or ratio <= 1 / 5:
                jump_found.append((c, qs[i - 1], round(prev, 1), qs[i], round(cur, 1), round(ratio, 3)))
    for c in NUMERIC_TOT_COLS:
        if c == "op_rate_total":
            continue
        s = tot.set_index("quarter")[c].sort_index()
        qs = list(s.index)
        for i in range(1, len(qs)):
            prev, cur = s.iloc[i - 1], s.iloc[i]
            if pd.isna(prev) or pd.isna(cur) or prev == 0 or cur == 0:
                continue
            ratio = cur / prev
            if ratio >= 5 or ratio <= 1 / 5:
                jump_found.append((c, qs[i - 1], round(prev, 1), qs[i], round(cur, 1), round(ratio, 3)))
    if jump_found:
        for row in jump_found:
            log(f"    지표={row[0]:<17} {row[1]}={row[2]:>12} -> {row[3]}={row[4]:>12}  배율={row[5]}")
    else:
        log("    없음 (2023Q4는 전업종 결측으로 합계가 0/NaN이 되어 비교에서 제외됨)")
    log(f"  [주의] 2025Q2 기계 production 스파이크 등은 값 제거·보간 대상이 아님 — 발견만 하고 그대로 둔다")
    record("D", f"단위 급변 후보 {len(jump_found)}건 발견 (검토용, 제거하지 않음)", True)


# ----------------------------------------------------------------------------
# E. 업종별 합계와 단지 전체의 관계
# ----------------------------------------------------------------------------

def check_e(ind: pd.DataFrame, tot: pd.DataFrame) -> pd.DataFrame:
    hr("E. 업종별 합계 vs 단지 전체")

    prod_sum = ind.groupby("quarter").production.sum(min_count=1)
    emp_sum = ind.groupby("quarter").employment.sum(min_count=1)
    t = tot.set_index("quarter")[["production_total", "employment_total"]]
    df = t.join(prod_sum.rename("prod_ind_sum")).join(emp_sum.rename("emp_ind_sum"))
    df = df.reset_index().rename(columns={"index": "quarter"})
    if "quarter" not in df.columns:
        df = df.rename(columns={df.columns[0]: "quarter"})

    df["prod_rel_err"] = np.where(
        df.production_total.notna() & df.prod_ind_sum.notna() & (df.production_total != 0),
        (df.prod_ind_sum - df.production_total).abs() / df.production_total.abs(),
        np.nan,
    )
    df["emp_gap"] = df.employment_total - df.emp_ind_sum

    df = df[["quarter", "prod_ind_sum", "production_total", "prod_rel_err",
             "emp_ind_sum", "employment_total", "emp_gap"]].sort_values("quarter")

    out_path = os.path.join(DIR_LOG, "total_vs_industry.csv")
    df.to_csv(out_path, index=False, encoding="utf-8-sig")
    log(f"  저장 -> {os.path.relpath(out_path, BASE)}")

    excl = df[df.quarter != "2023Q4"]
    max_err = excl.prod_rel_err.max()
    log(f"  production 상대오차 (2023Q4 제외) max={max_err:.6%} "
        f"median={excl.prod_rel_err.median():.6%}")
    record("E", f"업종합계 vs production_total 상대오차 <=1% (2023Q4 제외, 실제 max {max_err:.4%})",
           bool(max_err <= 0.01))

    log(f"  employment_total - 업종합계(nonmfg_implied, emp_gap) "
        f"min={df.emp_gap.min():.1f} max={df.emp_gap.max():.1f} median={df.emp_gap.median():.1f}")
    log(df[["quarter", "prod_rel_err", "emp_gap"]].tail(8).to_string(index=False))
    record("E", "업종별 employment 합계 vs employment_total 차이(emp_gap) 산출 완료", True)

    return df


# ----------------------------------------------------------------------------
# F. 연간보정본 비제조 고용 대조
# ----------------------------------------------------------------------------

def _period_key_to_quarter(key: str) -> tuple[str | None, bool]:
    """(quarter, is_quarter_end_month). Q키는 그대로 분기, M키는 속한 분기로 매핑하되
    분기말 월(3,6,9,12)인지 여부를 함께 반환한다 (고용은 stock이므로 그 달만 유효)."""
    if "Q" in key:
        return key, True
    if "M" in key:
        y, m = key.split("M")
        m = int(m)
        q = (m - 1) // 3 + 1
        return f"{y}Q{q}", m in (3, 6, 9, 12)
    return None, False


def check_f(gap_df: pd.DataFrame) -> None:
    hr("F. 연간보정본 비제조 고용 대조")

    if openpyxl is None:
        record("F", "openpyxl 미설치 -> 비제조 고용 대조 건너뜀", False)
        return

    paths = []
    for root, _d, fs in os.walk(DIR_REV):
        for f in fs:
            if re.match(r"^\d{4}[MQ]\d{1,2}\.xlsx$", f):
                paths.append(os.path.join(root, f))
    paths.sort()

    gap_lookup = gap_df.set_index("quarter")["emp_gap"].to_dict()

    recs = []
    for p in paths:
        key = os.path.splitext(os.path.basename(p))[0]
        q, is_qend = _period_key_to_quarter(key)
        if q is None:
            continue
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        sheet = next((s for s in wb.sheetnames if s.startswith("표9 ")), None)
        if sheet is None:
            wb.close()
            continue
        ws = wb[sheet]
        hdr = None
        cw_row = None
        for row in ws.iter_rows(values_only=True):
            first = str(row[0]).strip() if row and row[0] is not None else ""
            if first == "산업단지":
                hdr = [str(c).strip() if c is not None else "" for c in row]
            if first == "창원":
                cw_row = row
                break
        wb.close()
        if hdr is None or cw_row is None:
            continue

        def _val(colname):
            if colname in hdr:
                i = hdr.index(colname)
                if i < len(cw_row):
                    v = cw_row[i]
                    if v is None:
                        return None
                    try:
                        return float(str(v).replace(",", ""))
                    except ValueError:
                        return None
            return None

        mfg_sum = sum(v for v in (_val(ind) for ind in INDUSTRIES) if v is not None)
        nonmfg = _val("비제조")
        total_off = _val("계")

        applicable = is_qend  # stock 지표: 분기말 월만 유효
        emp_gap_master = gap_lookup.get(q)
        match = None
        if applicable and nonmfg is not None and emp_gap_master is not None and not pd.isna(emp_gap_master):
            match = abs(nonmfg - emp_gap_master) < 1.0  # 반올림 오차 허용 1명 미만

        recs.append(dict(revision_key=key, quarter=q, mfg_sum=mfg_sum,
                          nonmfg_official=nonmfg, total_official=total_off,
                          emp_gap_from_master=emp_gap_master,
                          applicable=applicable, match=match))

    df = pd.DataFrame(recs).sort_values("revision_key")
    out_path = os.path.join(DIR_LOG, "nonmfg_check.csv")
    df.drop(columns=["applicable"]).to_csv(out_path, index=False, encoding="utf-8-sig")
    log(f"  저장 -> {os.path.relpath(out_path, BASE)}")

    n_total = len(df)
    applicable_df = df[df.applicable]
    n_applicable = len(applicable_df)
    n_match = int((applicable_df.match == True).sum())
    n_mismatch = int((applicable_df.match == False).sum())
    diffs = (applicable_df.nonmfg_official - applicable_df.emp_gap_from_master).abs()
    max_diff = diffs.max() if len(diffs) else float("nan")

    log(f"  전체 revision 시점: {n_total}개")
    log(f"  분기말(stock 유효) 시점: {n_applicable}개")
    log(f"  일치: {n_match} / 불일치: {n_mismatch}")
    log(f"  최대 오차: {max_diff}")
    if n_mismatch:
        mm = applicable_df[applicable_df.match == False]
        log(mm[["revision_key", "quarter", "nonmfg_official", "emp_gap_from_master"]]
            .assign(diff=(mm.nonmfg_official - mm.emp_gap_from_master))
            .to_string(index=False))

    record("F", f"24개 시점 중 분기말 {n_applicable}개가 대조 대상, "
                f"그 중 {n_match}개 일치 / 최대오차 {max_diff}",
           n_applicable > 0 and n_mismatch == 0)


# ----------------------------------------------------------------------------
# G. 출처 추적
# ----------------------------------------------------------------------------

def check_g(ind: pd.DataFrame) -> None:
    hr("G. 출처 추적")

    na_counts = {c: int(ind[c].isna().sum()) for c in SOURCE_COLS}
    record("G", f"_source 컬럼 결측 없음 (실제 {na_counts})", all(v == 0 for v in na_counts.values()))

    bad_vals = {}
    for c in SOURCE_COLS:
        vals = set(ind[c].dropna().unique())
        bad = vals - ALLOWED_SOURCES
        if bad:
            bad_vals[c] = bad
    record("G", f"source 값이 허용 집합 안에 있음 (위반: {bad_vals or '없음'})", len(bad_vals) == 0)

    invalid_src_cols = [f"{k}_invalid_source" for k in
                        ["production", "employment", "op_rate", "firms_in", "firms_op"]]
    inv_counts = {c: int((ind[c] == 1).sum()) for c in invalid_src_cols if c in ind.columns}
    log(f"  _invalid_source==1 행 수: {inv_counts}")
    record("G", "_invalid_source==1 행 존재 여부 확인 완료 (정보성, 실패 아님)", True)


# ----------------------------------------------------------------------------
# H. YoY 재검증
# ----------------------------------------------------------------------------

def check_h(ind: pd.DataFrame) -> None:
    hr("H. YoY 재검증")

    py_ = ind.pivot(index="quarter", columns="industry", values="production").sort_index()
    ey_ = ind.pivot(index="quarter", columns="industry", values="employment").sort_index()
    chk_p = ((py_ / py_.shift(4) - 1) * 100).replace([np.inf, -np.inf], np.nan)
    chk_e = ((ey_ / ey_.shift(4) - 1) * 100).replace([np.inf, -np.inf], np.nan)

    stored_p = ind.pivot(index="quarter", columns="industry", values="production_yoy").sort_index()
    stored_e = ind.pivot(index="quarter", columns="industry", values="employment_yoy").sort_index()

    diff_p = (chk_p - stored_p).abs().to_numpy(dtype=float).ravel()
    diff_e = (chk_e - stored_e).abs().to_numpy(dtype=float).ravel()
    all_diffs = np.concatenate([diff_p, diff_e])
    all_diffs = all_diffs[~np.isnan(all_diffs)]
    max_err = float(all_diffs.max()) if len(all_diffs) else float("nan")
    log(f"  독립 재계산 vs 저장값 최대 절대오차: {max_err} (비교 가능한 셀 {len(all_diffs)}개)")
    record("H", f"YoY 독립 재계산 최대 절대오차 ~0 (실제 {max_err})",
           len(all_diffs) > 0 and max_err < 1e-6)

    for q in ["2023Q4", "2024Q4"]:
        n = int(ind.loc[ind.quarter == q, "production_yoy"].notna().sum())
        log(f"  {q} production_yoy 유효 업종 수 = {n}/10 (기대 0)")
        record("H", f"{q} production_yoy 전 업종 NaN (실제 유효 {n}/10)", n == 0)

    raw_vals = pd.concat([ind.production_yoy, ind.employment_yoy]).dropna().to_numpy(dtype=float)
    n_inf = int(np.isinf(raw_vals).sum())
    record("H", f"inf 잔존 0건 (실제 {n_inf})", n_inf == 0)


# ----------------------------------------------------------------------------
# main
# ----------------------------------------------------------------------------

def main() -> int:
    os.makedirs(DIR_LOG, exist_ok=True)
    log(f"QA 실행 {datetime.now():%Y-%m-%d %H:%M:%S} | base={os.path.basename(BASE)}")

    if not (os.path.exists(P_IND) and os.path.exists(P_TOT)):
        log(f"[오류] master 파일 없음: {P_IND} / {P_TOT}")
        with open(os.path.join(DIR_LOG, "qa_report.txt"), "w", encoding="utf-8") as f:
            f.write("\n".join(LOG_LINES))
        return 1

    ind = pd.read_csv(P_IND)
    tot = pd.read_csv(P_TOT)

    check_a(ind, tot)
    check_b(ind, tot)
    check_c(ind, tot)
    check_d(ind, tot)
    gap_df = check_e(ind, tot)
    check_f(gap_df)
    check_g(ind)
    check_h(ind)

    hr("요약")
    n_pass = sum(1 for _s, _l, p in RESULTS if p)
    n_fail = sum(1 for _s, _l, p in RESULTS if not p)
    log(f"총 {len(RESULTS)}개 검사 중 PASS {n_pass} / FAIL {n_fail}")
    if n_fail:
        log("\nFAIL 목록:")
        for s, l, p in RESULTS:
            if not p:
                log(f"  [{s}] {l}")

    with open(os.path.join(DIR_LOG, "qa_report.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(LOG_LINES))

    return 1 if n_fail else 0


if __name__ == "__main__":
    sys.exit(main())
